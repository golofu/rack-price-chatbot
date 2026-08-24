# -*- coding: utf-8 -*-
"""
สร้างไฟล์ Excel ตารางราคาแร็คพวงมาลัย 200 รายการ (ข้อมูลสมมติ สำหรับทดลองทำแชทบอท)
เลียนแบบรูปแบบตารางของ บริษัท อาร์เอ็มเอ ซิตี้ มอเตอร์ส จำกัด

รัน:  python build/seed_pricelist.py     (จาก root ของ repo — ใช้ครั้งแรกครั้งเดียว
                                          ถ้าจะแก้ข้อมูลจริงให้แก้ที่ data/*.xlsx โดยตรงแทน)

ออกไฟล์ (path สัมพัทธ์กับ repo ไม่ผูกกับเครื่องใดเครื่องหนึ่ง):
  data/rack-price-list-200.xlsx   ไฟล์ที่ build.py อ่านจริง
  data/rack-price-list-200.csv    UTF-8 BOM สำหรับป้อนบอทตัวอื่นหรือดูด้วย Excel
"""
import csv
import os
import random
import sys

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

random.seed(20260824)

HERE = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = os.path.join(os.path.dirname(HERE), "data")
BASE = "rack-price-list-200"

HYD = "น้ำมัน"
EPS = "ไฟฟ้ากระบอกแห้ง"

# (brand, prefix, [ (model, [ปีที่มีขาย...], [รหัสเครื่อง...], ประเภทแร็ค ) ])
CATALOG = [
    ("TOYOTA", "RT", [
        ("CAMRY SXV10",            ["1992-1996", "1994-1996"],              ["5S-FE", "3S-FE"], HYD),
        ("CAMRY SXV20",            ["1997-2000", "1999-2001"],              ["5S-FE", "1MZ-FE"], HYD),
        ("CAMRY ACV30",            ["2002-2006"],                            ["1AZ-FE", "2AZ-FE"], HYD),
        ("CAMRY ACV40",            ["2006-2012"],                            ["1AZ-FE", "2AZ-FE"], HYD),
        ("CAMRY ACV50",            ["2012-2018"],                            ["2AR-FE"], EPS),
        ("CAMRY ACV70",            ["2019-2023"],                            ["A25A-FKS"], EPS),
        ("COROLLA CROSS",          ["2020-2023"],                            ["2ZR-FXE"], EPS),
        ("RAV4",                   ["2013-2018"],                            ["3ZR-FE"], EPS),
        ("MAJESTY",                ["2019-2023"],                            ["2TR-FE"], HYD),
        ("HIACE LH184",            ["2001-2004"],                            ["5L", "3L"], HYD),
        ("HILUX TIGER",            ["1998-2004"],                            ["5L", "1KZ-TE"], HYD),
        ("SOLUNA VIOS",            ["2002-2007"],                            ["1NZ-FE"], HYD),
        ("VIGO 2WD",               ["2005-2011", "2011-2015"],               ["2KD-FTV", "1TR-FE", "2TR-FE"], HYD),
        ("VIGO 4WD",               ["2005-2011", "2011-2015"],               ["1KD-FTV", "2KD-FTV"], HYD),
        ("REVO 2WD",               ["2015-2020", "2020-2023"],               ["1GD-FTV", "2GD-FTV"], HYD),
        ("REVO 4WD",               ["2015-2020"],                            ["1GD-FTV"], HYD),
        ("INNOVA",                 ["2005-2011", "2011-2015", "2015-2020"],  ["1TR-FE", "2TR-FE", "2KD-FTV"], HYD),
        ("FORTUNER",               ["2005-2011", "2012-2015", "2015-2020", "2020-2023"], ["2KD-FTV", "1GD-FTV", "2TR-FE"], HYD),
        ("ALTIS 17.3 mm.",         ["2008-2012"],                            ["1ZR-FE", "2ZR-FE"], EPS),
        ("ALTIS 14 mm. (แกนเล็ก)", ["2008-2013"],                            ["1ZR-FE"], EPS),
        ("ALTIS",                  ["2001-2007", "2014-2019", "2019-2023"],  ["1ZZ-FE", "2ZR-FE"], EPS),
        ("VIOS",                   ["2003-2007", "2008-2012", "2013-2017", "2018-2022"], ["1NZ-FE", "2NR-FE"], EPS),
        ("YARIS",                  ["2006-2012", "2013-2017", "2017-2022"],  ["1NZ-FE", "2NR-FE"], EPS),
        ("YARIS ATIV",             ["2017-2022"],                            ["2NR-FE"], EPS),
        ("SIENTA",                 ["2016-2020"],                            ["2NR-FE"], EPS),
        ("AVANZA",                 ["2004-2011", "2012-2016", "2016-2021"],  ["K3-VE 1.3", "3SZ-VE 1.5"], EPS),
        ("COMMUTER KDH200 SHORT",  ["2005-2012"],                            ["2KD-FTV", "1TR-FE"], HYD),
        ("COMMUTER KDH220 LONG",   ["2005-2012", "2012-2018"],               ["2KD-FTV", "2TR-FE"], HYD),
        ("VENTURY",                ["2005-2012"],                            ["2TR-FE"], HYD),
        ("PRADO VZJ95",            ["1996-2001"],                            ["5VZ-FE", "1KZ-TE"], HYD),
        ("PRADO J120 / J100",      ["2003-2007"],                            ["1GR-FE", "5VZ-FE"], HYD),
        ("HARRIER G2",             ["2008-2012"],                            ["1MZ-FE 3.0", "2AZ-FE 2.4"], HYD),
        ("RX 300",                 ["2004-2008"],                            ["1MZ-FE 3.0", "2AZ-FE 2.4"], HYD),
        ("ALPHARD VELLFIRE",       ["2008-2014", "2015-2021"],               ["2GR-FE 3.5", "2AZ-FE 2.4"], EPS),
        ("ESTIMA (ไฮบริด)",         ["2006-2012"],                            ["2GR-FE 3.5", "2AZ-FE 2.4"], EPS),
        ("LH125 / LH112 รถตู้หัวจรวด", ["1992-2004"],                        ["2L", "3L", "5L", "2R"], HYD),
        ("AE100-AE111 (โฉมตองอ่อน)", ["1993-2000"],                          ["4A-FE", "5A-FE"], HYD),
        ("โซลูน่า (HI-TORQUE)",     ["1997-2002"],                            ["5A-FE"], HYD),
        ("HILUX MIGHTY-X",         ["1991-1998"],                            ["2L", "5L"], HYD),
        ("WISH",                   ["2003-2009"],                            ["1ZZ-FE", "2ZR-FE"], EPS),
        ("C-HR",                   ["2018-2022"],                            ["2ZR-FXE"], EPS),
    ]),
    ("HONDA", "RH", [
        ("CIVIC FD",       ["2006-2011"],              ["R18A", "K20Z"], EPS),
        ("CIVIC FB",       ["2012-2015"],              ["R18Z"], EPS),
        ("CIVIC FC",       ["2016-2021"],              ["L15B7 Turbo"], EPS),
        ("CIVIC DIMENSION", ["2001-2005"],             ["D17A"], HYD),
        ("ACCORD G7",      ["2003-2007"],              ["K24A", "J30A"], HYD),
        ("ACCORD G8",      ["2008-2012"],              ["K24Z", "J35Z"], HYD),
        ("ACCORD G9",      ["2013-2019"],              ["K24W"], EPS),
        ("CITY GM2",       ["2008-2013"],              ["L15A"], EPS),
        ("CITY GM6",       ["2014-2019"],              ["L15Z"], EPS),
        ("CITY GN2",       ["2020-2023"],              ["L15B Turbo"], EPS),
        ("JAZZ GE",        ["2008-2013"],              ["L15A"], EPS),
        ("JAZZ GK",        ["2014-2020"],              ["L15Z"], EPS),
        ("CR-V G2",        ["2002-2006"],              ["K24A1"], HYD),
        ("CR-V G3",        ["2007-2012"],              ["K24Z", "R20A"], HYD),
        ("CR-V G4",        ["2013-2016"],              ["R20A", "K24V"], EPS),
        ("HR-V",           ["2015-2021"],              ["L15Z"], EPS),
        ("BR-V",           ["2016-2021"],              ["L15Z"], EPS),
        ("MOBILIO",        ["2014-2020"],              ["L15Z"], EPS),
        ("STREAM",         ["2001-2006"],              ["K20A", "D17A"], HYD),
        ("FREED",          ["2008-2016"],              ["L15A"], EPS),
        ("ODYSSEY",        ["2004-2008"],              ["K24A"], HYD),
        ("CIVIC EK",       ["1996-2000"],              ["D15B", "D16Y"], HYD),
        ("ACCORD G6",      ["1998-2002"],              ["F20B", "F23A"], HYD),
        ("CITY TYPE Z",    ["1996-2002"],              ["D15B"], HYD),
        ("CR-V G1",        ["1996-2001"],              ["B20B"], HYD),
        ("JAZZ GD",        ["2003-2007"],              ["L13A", "L15A"], HYD),
    ]),
    ("NISSAN", "RN", [
        ("ALMERA",          ["2011-2019", "2020-2023"], ["HR15DE", "HR12DE Turbo"], EPS),
        ("MARCH",           ["2010-2019"],              ["HR12DE"], EPS),
        ("TIIDA",           ["2006-2012"],              ["HR16DE", "MR18DE"], EPS),
        ("SYLPHY",          ["2012-2019"],              ["MRA8DE"], EPS),
        ("TEANA J31",       ["2004-2008"],              ["QR25DE", "VQ23DE"], HYD),
        ("TEANA J32",       ["2009-2013"],              ["QR25DE", "VQ25DE"], HYD),
        ("X-TRAIL T30",     ["2001-2007"],              ["QR20DE", "QR25DE"], HYD),
        ("X-TRAIL T31",     ["2008-2014"],              ["MR20DE"], EPS),
        ("NAVARA D40",      ["2007-2014"],              ["YD25DDTi"], HYD),
        ("NAVARA NP300",    ["2015-2021"],              ["YD25DDTi"], HYD),
        ("FRONTIER D22",    ["1998-2006"],              ["QD32", "YD25"], HYD),
        ("SUNNY NEO",       ["2000-2006"],              ["QG16DE", "QG18DE"], HYD),
        ("CEFIRO A32",      ["1995-1999"],              ["VQ20DE"], HYD),
        ("CEFIRO A33",      ["2000-2004"],              ["VQ20DE", "VQ23DE"], HYD),
        ("URVAN NV350",     ["2013-2021"],              ["YD25DDTi", "QR25DE"], HYD),
        ("JUKE",            ["2014-2019"],              ["HR15DE"], EPS),
        ("TERRA",           ["2018-2022"],              ["YD25DDTi"], HYD),
        ("BIG-M BDI",       ["1993-1998"],              ["TD25", "TD27"], HYD),
        ("SUNNY B13",       ["1991-1995"],              ["GA16DE"], HYD),
        ("PULSAR",          ["2013-2016"],              ["MRA8DE"], EPS),
        ("NOTE",            ["2017-2021"],              ["HR12DE"], EPS),
        ("KICKS",           ["2020-2023"],              ["HR12DE e-POWER"], EPS),
        ("X-TRAIL T32",     ["2015-2020"],              ["MR20DD"], EPS),
    ]),
    ("ISUZU", "RI", [
        ("D-MAX 2WD",            ["2003-2011", "2012-2019", "2020-2023"], ["4JA1-L", "4JK1-TC", "4JJ1-TC", "RZ4E"], HYD),
        ("D-MAX 4WD",            ["2003-2011", "2012-2019", "2020-2023"], ["4JJ1-TC", "4JK1-TC"], HYD),
        ("D-MAX V-CROSS",        ["2012-2019", "2020-2023"],               ["4JJ1-TC", "RZ4E"], HYD),
        ("D-MAX X-SERIES",       ["2015-2019"],                            ["4JK1-TC"], HYD),
        ("FASTER",               ["1985-1991"],                            ["C223", "4JA1"], HYD),
        ("TROOPER",              ["1992-2002"],                            ["4JG2", "4JX1"], HYD),
        ("D-MAX HI-LANDER",      ["2007-2011", "2012-2019"],              ["4JK1-TC"], HYD),
        ("MU-7",                 ["2004-2013"],                            ["4JJ1-TC", "4JK1-TC"], HYD),
        ("MU-X",                 ["2014-2019", "2020-2023"],               ["4JK1-TC", "RZ4E"], HYD),
        ("TFR มังกรทอง",          ["1991-1997"],                            ["4JA1", "4JB1-T"], HYD),
        ("DRAGON EYES",          ["1998-2002"],                            ["4JA1-L", "4JB1-T"], HYD),
        ("ADVENTURE MASTER",     ["1999-2002"],                            ["4JA1-L"], HYD),
        ("SPACECAB",             ["2003-2011"],                            ["4JA1-L"], HYD),
    ]),
    ("MITSUBISHI", "RM", [
        ("TRITON 2WD",     ["2005-2014", "2015-2023"], ["4D56", "4N15"], HYD),
        ("TRITON 4WD",     ["2005-2014", "2015-2023"], ["4D56", "4N15"], HYD),
        ("PAJERO SPORT",   ["2008-2015", "2016-2022"], ["4D56", "4N15"], HYD),
        ("LANCER CEDIA",   ["2001-2006"],              ["4G93", "4G94"], HYD),
        ("LANCER EX",      ["2009-2015"],              ["4B10", "4B11"], EPS),
        ("MIRAGE",         ["2012-2019"],              ["3A92"], EPS),
        ("ATTRAGE",        ["2013-2019", "2020-2023"], ["3A92"], EPS),
        ("XPANDER",        ["2018-2023"],              ["4A91"], EPS),
        ("STRADA",         ["1996-2005"],              ["4D56", "4G64"], HYD),
        ("SPACE WAGON",    ["1997-2003", "2004-2011"], ["4G63", "4G69"], HYD),
        ("TRITON ATHLETE", ["2019-2023"],              ["4N15"], HYD),
        ("LANCER GLXi",    ["1992-1996"],              ["4G13", "4G92"], HYD),
        ("GALANT",         ["1993-1998"],              ["4G63", "6A12"], HYD),
    ]),
    ("MAZDA", "RZ", [
        ("MAZDA 2 (SEDAN)",   ["2009-2014", "2015-2022"], ["ZY-VE", "P3 SkyActiv"], EPS),
        ("MAZDA 2 (HATCH)",   ["2009-2014", "2015-2022"], ["ZY-VE", "P3 SkyActiv"], EPS),
        ("MAZDA 3 BK",        ["2005-2010"],              ["LF-VE", "ZY-VE"], HYD),
        ("MAZDA 3 BL",        ["2011-2014"],              ["LF-VE"], EPS),
        ("MAZDA 3 BM",        ["2014-2019"],              ["PE SkyActiv"], EPS),
        ("MAZDA 6",           ["2003-2008"],              ["L3-VE"], HYD),
        ("CX-3",              ["2016-2021"],              ["P5 SkyActiv"], EPS),
        ("CX-5",              ["2013-2017", "2018-2022"], ["PE SkyActiv", "SH SkyActiv-D"], EPS),
        ("BT-50",             ["2007-2011"],              ["WL-C", "WE-C"], HYD),
        ("BT-50 PRO",         ["2012-2020"],              ["WE-C", "P4AT"], HYD),
        ("MAZDA 323 PROTEGE", ["1998-2004"],              ["FP-DE", "ZM-DE"], HYD),
        ("MAZDA 3 BP",        ["2019-2023"],              ["PE SkyActiv-G"], EPS),
        ("CX-30",             ["2020-2023"],              ["PE SkyActiv-G"], EPS),
        ("CX-8",              ["2019-2023"],              ["SH SkyActiv-D"], EPS),
        ("FIGHTER",           ["1998-2006"],              ["WL-T", "WL"], HYD),
    ]),
    ("FORD", "RF", [
        ("RANGER 2WD",     ["2006-2011", "2012-2018", "2019-2022"], ["WLC 2.5", "P4AT 2.2", "P5AT 3.2"], HYD),
        ("RANGER 4WD",     ["2006-2011", "2012-2018", "2019-2022"], ["WEC 3.0", "P5AT 3.2"], HYD),
        ("EVEREST",        ["2007-2014", "2015-2021", "2022-2023"], ["WEC 3.0", "P4AT 2.2"], HYD),
        ("COURIER",        ["1998-2005"],                            ["WL 2.5"], HYD),
        ("FIESTA",         ["2011-2016"],                            ["Duratec 1.5", "EcoBoost 1.0"], EPS),
        ("FOCUS",          ["2005-2011", "2012-2018"],               ["Duratec 1.8", "Duratec 2.0"], EPS),
        ("ECOSPORT",       ["2014-2019"],                            ["Duratec 1.5"], EPS),
    ]),
    ("CHEVROLET", "RC", [
        ("COLORADO 2WD",  ["2004-2011", "2012-2019"], ["4JA1-L", "Duramax 2.5", "Duramax 2.8"], HYD),
        ("COLORADO 4WD",  ["2004-2011", "2012-2019"], ["Duramax 2.8"], HYD),
        ("CAPTIVA",       ["2007-2012", "2013-2018"], ["Z20S1 2.0 D", "LNQ 2.4"], HYD),
        ("TRAILBLAZER",   ["2013-2019"],              ["Duramax 2.8"], HYD),
        ("AVEO",          ["2006-2012"],              ["F16D3"], EPS),
        ("CRUZE",         ["2011-2016"],              ["F18D4"], EPS),
    ]),
    ("SUZUKI", "RS", [
        ("SWIFT",   ["2012-2017", "2018-2023"], ["K12B", "K12M"], EPS),
        ("CIAZ",    ["2016-2022"],              ["K12M"], EPS),
        ("ERTIGA",  ["2013-2018", "2019-2023"], ["K14B", "K15B"], EPS),
        ("CARRY",   ["2014-2022"],              ["K15B"], HYD),
    ]),
    ("MG", "RG", [
        ("MG3",     ["2015-2021"],  ["1.5 VTi"], EPS),
        ("MG ZS",   ["2018-2023"],  ["1.5 VTi", "1.0T"], EPS),
        ("MG5",     ["2021-2023"],  ["1.5 VTi"], EPS),
        ("MG HS",   ["2019-2023"],  ["1.5T"], EPS),
        ("MG EXTENDER", ["2019-2023"], ["2.0 Turbo"], HYD),
    ]),
]

TARGET_ROWS = 200


def price_for(rack_type, lhd):
    """สุ่มราคาให้สมเหตุผลตามประเภทแร็ค ปัดเป็นหลักร้อย"""
    if rack_type == HYD:
        base = random.randrange(3200, 5001, 100)
    else:
        base = random.randrange(2300, 3401, 100)
    if lhd:                       # พวงมาลัยซ้าย หายากกว่า บวกราคา
        base += random.choice([500, 600, 700])
    return base


def build_rows():
    rows = []
    # โควตาจำนวนแถวต่อยี่ห้อ ให้ TOYOTA เยอะสุดเหมือนต้นฉบับ
    quota = {"TOYOTA": 62, "HONDA": 26, "NISSAN": 24, "ISUZU": 20, "MITSUBISHI": 18,
             "MAZDA": 18, "FORD": 14, "CHEVROLET": 8, "SUZUKI": 6, "MG": 4}

    for brand, prefix, models in CATALOG:
        want = quota[brand]
        # แตกทุก (model, year) เป็นตัวเลือก แล้วสุ่มเลือกจนครบโควตา
        variants = []
        for model, years, engines, rtype in models:
            for y in years:
                variants.append((model, y, engines, rtype))
        assert len(variants) >= want, "%s มีรุ่นให้เลือกแค่ %d ต้องการ %d" % (
            brand, len(variants), want)

        # เลือกแบบไม่ซ้ำ รุ่น+ปี จะไม่โผล่สองครั้งในยี่ห้อเดียวกัน
        picked = random.sample(variants, want)

        # จับกลุ่มเฉพาะรถที่ใช้แร็คประเภทเดียวกัน (รหัสเดียวห้ามคละน้ำมัน/ไฟฟ้า)
        by_type = {HYD: [], EPS: []}
        for v in picked:
            by_type[v[3]].append(v)

        code_no = 1
        for rtype in (HYD, EPS):
            bucket = by_type[rtype]
            random.shuffle(bucket)
            idx = 0
            while idx < len(bucket):
                group = random.choice([1, 2, 2, 3])
                lhd = random.random() < 0.18      # ~18% เป็นรุ่นพวงมาลัยซ้าย
                code = "%s%s%02d" % (prefix, "L" if lhd else "", code_no)
                for model, year, engines, _rt in bucket[idx:idx + group]:
                    note = " *พวงมาลัยซ้าย*" if lhd else ""
                    rows.append({
                        "code": code,
                        "brand": brand,
                        "model": model + note,
                        "year": year,
                        "engine": random.choice(engines),
                        "type": rtype,
                        "price": price_for(rtype, lhd),
                    })
                idx += group
                code_no += 1

    assert len(rows) == TARGET_ROWS, "ได้ %d แถว ไม่ใช่ %d" % (len(rows), TARGET_ROWS)

    # ราคาเดียวกันทั้งรหัสเดียวกัน (ของจริงรหัสเดียวราคาเดียว)
    price_by_code = {}
    for r in rows:
        key = (r["brand"], r["code"])
        price_by_code.setdefault(key, r["price"])
        r["price"] = price_by_code[key]

    for n, r in enumerate(rows, start=1):
        r["no"] = n
    return rows


HEADERS = ["ลำดับ", "รหัสแร็ค", "BRAND", "MODEL", "ปี", "เครื่องยนต์",
           "ประเภทแร็ค", "ราคาขาย (ร้านอะไหล่)"]
WIDTHS = [7, 12, 14, 34, 13, 18, 18, 20]


def write_xlsx(rows, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "ราคาแร็ค"

    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill("solid", fgColor="D9D9D9")

    ws.append(HEADERS)
    for c in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(name="Tahoma", size=10, bold=True)
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        ws.column_dimensions[get_column_letter(c)].width = WIDTHS[c - 1]
    ws.row_dimensions[1].height = 30

    for r in rows:
        ws.append([r["no"], r["code"], r["brand"], r["model"], r["year"],
                   r["engine"], r["type"], r["price"]])

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(HEADERS)):
        for cell in row:
            cell.font = Font(name="Tahoma", size=10)
            cell.border = border
            if cell.column in (1, 2, 3, 5, 7):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif cell.column == 8:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = "#,##0"
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(HEADERS)), ws.max_row)
    wb.save(path)


def write_csv(rows, path):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(HEADERS)
        for r in rows:
            w.writerow([r["no"], r["code"], r["brand"], r["model"], r["year"],
                        r["engine"], r["type"], r["price"]])


if __name__ == "__main__":
    rows = build_rows()

    # ---- ตรวจความสมเหตุสมผลของข้อมูล "ก่อน" เขียนไฟล์ลงดิสก์เสมอ ----
    # เดิมบล็อกนี้อยู่ "หลัง" write_xlsx/write_csv แปลว่าถ้าข้อมูลพัง ไฟล์ที่คละราคา/
    # ประเภทต่อรหัสก็เขียนลงดิสก์ไปแล้วก่อน assert จะทำงาน (build.py อาจหยิบไฟล์เสียไปใช้)
    type_of_code, price_of_code, seen = {}, {}, set()
    for r in rows:
        key = (r["brand"], r["code"])
        assert type_of_code.setdefault(key, r["type"]) == r["type"], \
            "รหัส %s คละประเภทแร็ค" % r["code"]
        assert price_of_code.setdefault(key, r["price"]) == r["price"], \
            "รหัส %s ราคาไม่ตรงกัน" % r["code"]
        mk = (r["brand"], r["model"], r["year"])
        assert mk not in seen, "รุ่นซ้ำ: %s" % (mk,)
        seen.add(mk)

    os.makedirs(OUT_DIR, exist_ok=True)
    xlsx_path = os.path.join(OUT_DIR, BASE + ".xlsx")
    csv_path = os.path.join(OUT_DIR, BASE + ".csv")

    # data/*.xlsx คือไฟล์ที่ build.py อ่านจริง ถ้ามีคนแก้ราคาจริงไว้แล้วรันสคริปต์นี้ซ้ำ
    # จะเขียนทับข้อมูลสมมติกลับไปทันทีโดยไม่มีการเตือน จึงต้องกัน --force ไว้
    if os.path.exists(xlsx_path) and "--force" not in sys.argv:
        raise SystemExit(
            "%s มีอยู่แล้ว สคริปต์นี้ทำไว้ seed ข้อมูลตัวอย่างครั้งแรกเท่านั้น\n"
            "ถ้าตั้งใจจะเขียนทับ (ราคาจริงที่แก้ไว้จะหาย) ให้รันใหม่พร้อม --force" % xlsx_path)

    write_xlsx(rows, xlsx_path)
    write_csv(rows, csv_path)

    brands = {}
    for r in rows:
        brands[r["brand"]] = brands.get(r["brand"], 0) + 1
    print("rows:", len(rows))
    print("brands:", brands)
    print("codes:", len(set((r["brand"], r["code"]) for r in rows)))
    print("saved:", xlsx_path)
    print("saved:", csv_path)
    print("ถัดไป: python build/build.py  เพื่อประกอบ index.html ใหม่")
