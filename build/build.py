# -*- coding: utf-8 -*-
"""ประกอบ index.html: template + ข้อมูลจาก data/rack-price-list-200.xlsx + โลโก้ Gen AI Space

รัน:  python build/build.py        (จาก root ของ repo)
ต้องมี: pip install openpyxl pillow
"""
import base64
import io
import json
import os
import re

import openpyxl
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)

XLSX = os.path.join(ROOT, "data", "rack-price-list-200.xlsx")
LOGO = os.path.join(HERE, "logo-gen-ai-space.png")
TPL = os.path.join(HERE, "template.html")
OUT = os.path.join(ROOT, "index.html")

TYPES = ["น้ำมัน", "ไฟฟ้ากระบอกแห้ง"]
LHD_MARK = "*พวงมาลัยซ้าย*"
HEADERS = ["ลำดับ", "รหัสแร็ค", "BRAND", "MODEL", "ปี", "เครื่องยนต์",
           "ประเภทแร็ค", "ราคาขาย (ร้านอะไหล่)"]


def load_rows():
    ws = openpyxl.load_workbook(XLSX).active

    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    if header != HEADERS:
        raise SystemExit(
            "หัวตารางไม่ตรงกับที่ระบบต้องการ\n  ต้องเป็น: %s\n  แต่เจอ:   %s"
            % (HEADERS, header))

    brands, out = [], []
    skipped = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        no, code, brand, model, year, engine, rtype, price = row[:8]
        if code is None:
            # แถวว่างจริง (ไม่มีข้อมูลอะไรเลย) ข้ามแบบเงียบได้ แต่ถ้ามีข้อมูลอื่นอยู่
            # แสดงว่าพิมพ์ตกช่องรหัสแร็ค ต้องเตือน ไม่งั้นรถหายไปจากเว็บโดยไม่มีใครรู้
            if any(v is not None for v in (brand, model, year, engine, rtype, price)):
                skipped.append(i)
            continue
        if not brand:
            raise SystemExit('แถวที่ %d: ไม่มี BRAND (รหัสแร็ค %r)' % (i, code))
        if not model:
            raise SystemExit('แถวที่ %d: ไม่มี MODEL (รหัสแร็ค %r)' % (i, code))
        if rtype not in TYPES:
            raise SystemExit(
                'แถวที่ %d: ประเภทแร็ค "%s" ไม่ถูกต้อง ใช้ได้แค่ %s' % (i, rtype, TYPES))
        if not re.fullmatch(r"\d{4}-\d{4}", str(year or "")):
            raise SystemExit('แถวที่ %d: ปี "%s" ต้องอยู่ในรูป YYYY-YYYY' % (i, year))
        price_str = str(price).replace(",", "").strip()
        try:
            price_int = int(price_str)
        except (TypeError, ValueError):
            try:
                float(price_str)
            except (TypeError, ValueError):
                raise SystemExit(
                    'แถวที่ %d: ราคา "%s" ไม่ใช่ตัวเลข (รหัสแร็ค %r)' % (i, price, code))
            raise SystemExit(
                'แถวที่ %d: ราคา "%s" มีจุดทศนิยม ระบบรับเฉพาะราคาเต็มบาท (รหัสแร็ค %r)'
                % (i, price, code))
        if brand not in brands:
            brands.append(brand)
        lhd = 1 if LHD_MARK in str(model) else 0
        clean = re.sub(r"\s*\*พวงมาลัยซ้าย\*", "", str(model)).strip()
        out.append([code, brands.index(brand), clean, str(year), engine,
                    TYPES.index(rtype), price_int, lhd])
    if skipped:
        raise SystemExit(
            "แถวที่ %s: มีข้อมูลอยู่แต่ช่องรหัสแร็คว่างเปล่า — ตรวจว่าพิมพ์ตกหรือลบผิดช่อง"
            % ", ".join(str(n) for n in skipped))
    if not out:
        raise SystemExit("ไม่พบข้อมูลในไฟล์ xlsx")

    # หน้าเว็บถือว่า 1 รหัสแร็ค = 1 ราคา 1 ประเภท ถ้าข้อมูลขัดกันการ์ดจะโชว์ราคาผิด
    seen = {}
    for code, bi, model, year, engine, ti, price, lhd in out:
        key = (brands[bi], code)
        if key in seen:
            prev_price, prev_type = seen[key]
            if prev_price != price:
                raise SystemExit(
                    "รหัส %s (%s) มีสองราคา: %s กับ %s — รหัสเดียวต้องราคาเดียว"
                    % (code, brands[bi], prev_price, price))
            if prev_type != ti:
                raise SystemExit(
                    "รหัส %s (%s) คละประเภทแร็ค: %s กับ %s"
                    % (code, brands[bi], TYPES[prev_type], TYPES[ti]))
        else:
            seen[key] = (price, ti)

    return brands, out


def logo_uri():
    im = Image.open(LOGO).convert("RGB").resize((192, 192), Image.LANCZOS)
    buf = io.BytesIO()
    im.save(buf, "WEBP", quality=88, method=6)
    return "data:image/webp;base64," + base64.b64encode(buf.getvalue()).decode()


def main():
    brands, rows = load_rows()
    db = json.dumps({"b": brands, "r": rows}, ensure_ascii=False, separators=(",", ":"))

    with open(TPL, encoding="utf-8") as f:
        html = f.read()
    if "/*__DATA__*/" not in html or "__LOGO_URI__" not in html:
        raise SystemExit("template.html ไม่มี placeholder /*__DATA__*/ หรือ __LOGO_URI__")
    html = html.replace("/*__DATA__*/", db).replace("__LOGO_URI__", logo_uri())

    with open(OUT, "w", encoding="utf-8") as f:
        f.write(html)

    print("rows: %d  brands: %d" % (len(rows), len(brands)))
    print("size: %.1f KB" % (len(html.encode("utf-8")) / 1024))
    print("built:", OUT)


if __name__ == "__main__":
    main()
